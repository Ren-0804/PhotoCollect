import tempfile
import threading
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from report_services import (
    ReportExportCancelled,
    build_collection_report,
    export_collection_report_xlsx,
)


CATEGORIES = [
    (1, "1 全开门空箱"),
    (2, "2 全开门1/2箱"),
    (3, "3 全开门满箱"),
    (4, "4 半关门满箱"),
    (5, "5 全关门"),
    (6, "6 加固/底梁 1"),
    (7, "7 加固/底梁 2"),
    (8, "8 铅封/其他"),
]


def sample_summaries():
    return [
        {
            "name": "MSCU6639870",
            "ocr_size_type": "22G1",
            "task_state": "complete",
            "ocr_status": "match",
            "pic_cnt": 2,
            "required_categories": [1, 2],
            "category_counts": {"1": 1, "2": 1},
            "last_updated": "2026-07-11 09:00:00",
        },
        {
            "name": "LYGU3043781",
            "ocr_size_type": "22G1",
            "task_state": "pending",
            "ocr_status": "",
            "pic_cnt": 1,
            "required_categories": [1, 2, 3],
            "category_counts": {"1": 1, "2": 0, "3": 0},
            "last_updated": "2026-07-11 10:00:00",
        },
        {
            "name": "INVALID-TASK",
            "ocr_size_type": "",
            "task_state": "anomaly",
            "ocr_status": "expected_invalid",
            "pic_cnt": 1,
            "required_categories": [1],
            "category_counts": {"1": 1},
            "last_updated": "2026-07-11 11:00:00",
        },
    ]


class CollectionReportTests(unittest.TestCase):
    def test_report_respects_per_task_requirements(self):
        report = build_collection_report(sample_summaries(), CATEGORIES)

        self.assertEqual(report.total_tasks, 3)
        self.assertEqual(report.complete_tasks, 1)
        self.assertEqual(report.pending_tasks, 1)
        self.assertEqual(report.anomaly_tasks, 1)
        self.assertEqual(report.total_photos, 4)
        self.assertEqual(report.photo_complete_tasks, 2)
        self.assertEqual(report.missing_photo_tasks, 1)
        self.assertEqual(report.verification_issue_tasks, 1)
        self.assertEqual(report.required_slots, 6)
        self.assertEqual(report.collected_required_slots, 4)
        self.assertAlmostEqual(report.photo_completion_rate, 2 / 3)
        self.assertEqual(len(report.issue_rows), 2)

        category_1, category_2, category_3 = report.categories[:3]
        self.assertEqual((category_1.required_tasks, category_1.collected_tasks, category_1.missing_tasks), (3, 3, 0))
        self.assertEqual((category_2.required_tasks, category_2.collected_tasks, category_2.missing_tasks), (2, 1, 1))
        self.assertEqual((category_3.required_tasks, category_3.collected_tasks, category_3.missing_tasks), (1, 0, 1))

        pending = next(row for row in report.rows if row.name == "LYGU3043781")
        self.assertEqual(pending.missing_category_ids, (2, 3))
        self.assertEqual(pending.missing_labels, ("全开门1/2箱", "全开门满箱"))

    def test_non_contiguous_category_ids_keep_counts_aligned(self):
        report = build_collection_report(
            [{
                "name": "TEST",
                "required_categories": [3],
                "category_counts": {"3": 2},
                "pic_cnt": 2,
            }],
            [(1, "1 第一类"), (3, "3 第三类")],
        )

        row = report.rows[0]
        self.assertEqual(row.category_count(1), 0)
        self.assertEqual(row.category_count(3), 2)
        self.assertEqual(row.missing_category_ids, ())
        self.assertEqual(report.categories[1].uploaded_photos, 2)

    def test_malformed_counts_are_treated_as_zero(self):
        report = build_collection_report(
            [{
                "name": "TEST",
                "required_categories": [1, 2, 3, 4],
                "category_counts": {
                    "1": "bad",
                    "2": None,
                    "3": -4,
                    "4": float("inf"),
                },
                "pic_cnt": "bad",
            }],
            CATEGORIES,
        )

        row = report.rows[0]
        self.assertEqual(tuple(row.category_count(index) for index in range(1, 5)), (0, 0, 0, 0))
        self.assertEqual(row.missing_category_ids, (1, 2, 3, 4))
        self.assertEqual(row.photo_count, 0)
        self.assertEqual(report.total_photos, 0)

    def test_excel_contains_dashboard_details_and_photo_matrix(self):
        report = build_collection_report(sample_summaries(), CATEGORIES)
        with tempfile.TemporaryDirectory() as tmp:
            destination = Path(tmp) / "collection-report.xlsx"

            output = export_collection_report_xlsx(
                report,
                destination,
                Path(tmp) / "work-root",
                generated_at=datetime(2026, 7, 11, 12, 30),
            )

            self.assertEqual(output, destination)
            self.assertTrue(output.is_file())
            self.assertEqual(list(Path(tmp).glob("*.partial.xlsx")), [])
            workbook = load_workbook(output, data_only=False)
            self.assertEqual(workbook.sheetnames, ["采集看板", "任务明细", "照片矩阵"])

            dashboard = workbook["采集看板"]
            self.assertEqual(dashboard["A1"].value, "PhotoCollect 采集看板")
            self.assertEqual(dashboard["B5"].value, 3)
            self.assertEqual(dashboard["D5"].value, 2)
            self.assertEqual(dashboard["F5"].value, 1)
            self.assertEqual(dashboard["B6"].value, 1)
            self.assertEqual(dashboard["D6"].value, 4)
            self.assertEqual(dashboard["F6"].value, "=IF(SUM(C10:C17)=0,0,SUM(D10:D17)/SUM(C10:C17))")
            self.assertEqual(dashboard["G6"].value, "=IF(B5=0,0,D5/B5)")
            self.assertEqual(dashboard["E11"].value, "=C11-D11")
            self.assertEqual(dashboard["F11"].value, '=IF(C11=0,"-",D11/C11)')
            self.assertEqual(dashboard["F14"].value, '=IF(C14=0,"-",D14/C14)')

            details = workbook["任务明细"]
            self.assertEqual(details["A2"].value, "LYGU3043781")
            self.assertEqual(details["C2"].value, "缺照片")
            self.assertIn("全开门1/2箱", details["G2"].value)
            self.assertEqual(details["A3"].value, "INVALID-TASK")
            self.assertEqual(details["C3"].value, "照片齐全")
            self.assertEqual(details["D3"].value, "任务箱号异常")

            matrix = workbook["照片矩阵"]
            self.assertEqual(matrix["A2"].value, "LYGU3043781")
            self.assertEqual(matrix["C2"].value, "缺照片")
            self.assertEqual(matrix["E2"].value, "缺失")
            self.assertEqual(matrix["F2"].value, "缺失")
            workbook.close()

    def test_failed_excel_export_removes_partial_file(self):
        report = build_collection_report(sample_summaries(), CATEGORIES)
        with tempfile.TemporaryDirectory() as tmp:
            destination = Path(tmp) / "collection-report.xlsx"
            with patch("report_services.Workbook.save", side_effect=OSError("disk full")):
                with self.assertRaises(OSError):
                    export_collection_report_xlsx(report, destination, tmp)

            self.assertFalse(destination.exists())
            self.assertEqual(list(Path(tmp).glob("*.partial.xlsx")), [])

    def test_excel_export_reports_phases_and_cancels_before_publication(self):
        report = build_collection_report(sample_summaries(), CATEGORIES)
        with tempfile.TemporaryDirectory() as tmp:
            destination = Path(tmp) / "collection-report.xlsx"
            destination.write_bytes(b"previous-report")
            cancel_event = threading.Event()
            phases = []

            def on_progress(current, total, label):
                phases.append((current, total, label))
                if current == 4:
                    cancel_event.set()

            with self.assertRaises(ReportExportCancelled):
                export_collection_report_xlsx(
                    report,
                    destination,
                    tmp,
                    cancel_event=cancel_event,
                    progress_callback=on_progress,
                )

            self.assertEqual(destination.read_bytes(), b"previous-report")
            self.assertEqual([current for current, _total, _label in phases], [1, 2, 3, 4])
            self.assertTrue(all(total == 5 for _current, total, _label in phases))
            self.assertEqual(list(Path(tmp).glob("*.partial.xlsx")), [])


if __name__ == "__main__":
    unittest.main()
