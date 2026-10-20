"""Collection reporting and Excel export for PhotoCollect.

This module only consumes task summaries. It never mutates photo folders or
metadata, so report generation stays independent from collection workflows.
"""

from __future__ import annotations

import os
import re
import uuid
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


TASK_STATUS_LABELS = {
    "complete": "已完成",
    "pending": "待补照片",
    "anomaly": "异常",
}

OCR_STATUS_LABELS = {
    "match": "已核对",
    "inferred_match": "校验位推算吻合，需确认",
    "mismatch": "箱号不一致",
    "unreadable": "未识别",
    "expected_invalid": "任务箱号异常",
    "error": "OCR失败",
    "": "待识别",
}


class ReportExportCancelled(Exception):
    """Raised when a caller cancels an Excel report before publication."""


@dataclass(frozen=True)
class ReportRow:
    name: str
    size_type: str
    task_state: str
    ocr_status: str
    photo_count: int
    required_categories: tuple[int, ...]
    category_counts: tuple[int, ...]
    missing_category_ids: tuple[int, ...]
    missing_labels: tuple[str, ...]
    last_updated: str

    def category_count(self, category: int) -> int:
        if 1 <= category <= len(self.category_counts):
            return self.category_counts[category - 1]
        return 0

    @property
    def photo_complete(self) -> bool:
        return not self.missing_category_ids

    @property
    def verification_issue(self) -> bool:
        return self.task_state == "anomaly"


@dataclass(frozen=True)
class CategoryReport:
    index: int
    label: str
    required_tasks: int
    collected_tasks: int
    missing_tasks: int
    uploaded_photos: int

    @property
    def completion_rate(self) -> float:
        return self.collected_tasks / self.required_tasks if self.required_tasks else 0.0


@dataclass(frozen=True)
class CollectionReport:
    rows: tuple[ReportRow, ...]
    categories: tuple[CategoryReport, ...]
    total_tasks: int
    complete_tasks: int
    pending_tasks: int
    anomaly_tasks: int
    photo_complete_tasks: int
    missing_photo_tasks: int
    verification_issue_tasks: int
    total_photos: int
    required_slots: int
    collected_required_slots: int

    @property
    def task_completion_rate(self) -> float:
        return self.complete_tasks / self.total_tasks if self.total_tasks else 0.0

    @property
    def photo_task_completion_rate(self) -> float:
        return self.photo_complete_tasks / self.total_tasks if self.total_tasks else 0.0

    @property
    def photo_completion_rate(self) -> float:
        return self.collected_required_slots / self.required_slots if self.required_slots else 0.0

    @property
    def issue_rows(self) -> tuple[ReportRow, ...]:
        return tuple(row for row in self.rows if not row.photo_complete or row.verification_issue)


def _clean_category_label(label: str) -> str:
    return re.sub(r"^\s*\d+\s*", "", str(label)).strip()


def _coerce_nonnegative_int(value) -> int:
    try:
        return max(0, int(value or 0))
    except (TypeError, ValueError, OverflowError):
        return 0


def build_collection_report(task_summaries: Sequence[Mapping], categories) -> CollectionReport:
    """Build report metrics from dashboard task summaries."""
    normalized_categories = []
    seen_categories = set()
    for raw_index, label in categories:
        try:
            index = int(raw_index)
        except (TypeError, ValueError, OverflowError):
            continue
        if index <= 0 or index in seen_categories:
            continue
        seen_categories.add(index)
        normalized_categories.append((index, _clean_category_label(label)))
    category_pairs = tuple(normalized_categories)
    category_labels = dict(category_pairs)
    max_category_index = max(category_labels, default=0)
    rows = []

    for summary in task_summaries:
        required = tuple(sorted({
            int(value) for value in summary.get("required_categories", ())
            if str(value).isdigit() and int(value) in category_labels
        }))
        raw_counts = summary.get("category_counts", {})
        counts_by_index = [0] * max_category_index
        if isinstance(raw_counts, Mapping):
            for index, _label in category_pairs:
                counts_by_index[index - 1] = _coerce_nonnegative_int(
                    raw_counts.get(str(index), raw_counts.get(index, 0))
                )
        counts = tuple(counts_by_index)
        raw_missing_ids = summary.get("missing_category_ids")
        if isinstance(raw_missing_ids, (list, tuple, set)):
            missing_ids = tuple(sorted({
                int(value)
                for value in raw_missing_ids
                if str(value).isdigit() and int(value) in required
            }))
        else:
            missing_ids = tuple(index for index in required if counts[index - 1] == 0)
        missing_labels = tuple(category_labels[index] for index in missing_ids)
        rows.append(ReportRow(
            name=str(summary.get("name", "")),
            size_type=str(summary.get("ocr_size_type", "")),
            task_state=str(summary.get("task_state", "pending")),
            ocr_status=str(summary.get("ocr_status", "")),
            photo_count=_coerce_nonnegative_int(summary.get("pic_cnt", 0)),
            required_categories=required,
            category_counts=counts,
            missing_category_ids=missing_ids,
            missing_labels=missing_labels,
            last_updated=str(summary.get("last_updated", "")),
        ))

    rows.sort(key=lambda row: (
        0 if not row.photo_complete else 1 if row.verification_issue else 2,
        row.name,
    ))
    category_reports = []
    for index, label in category_pairs:
        required_rows = [row for row in rows if index in row.required_categories]
        collected = sum(1 for row in required_rows if row.category_count(index) > 0)
        category_reports.append(CategoryReport(
            index=index,
            label=label,
            required_tasks=len(required_rows),
            collected_tasks=collected,
            missing_tasks=len(required_rows) - collected,
            uploaded_photos=sum(row.category_count(index) for row in rows),
        ))

    complete = sum(row.task_state == "complete" for row in rows)
    pending = sum(row.task_state == "pending" for row in rows)
    anomaly = sum(row.task_state == "anomaly" for row in rows)
    required_slots = sum(len(row.required_categories) for row in rows)
    collected_slots = sum(
        row.category_count(index) > 0
        for row in rows
        for index in row.required_categories
    )
    return CollectionReport(
        rows=tuple(rows),
        categories=tuple(category_reports),
        total_tasks=len(rows),
        complete_tasks=complete,
        pending_tasks=pending,
        anomaly_tasks=anomaly,
        photo_complete_tasks=sum(row.photo_complete for row in rows),
        missing_photo_tasks=sum(not row.photo_complete for row in rows),
        verification_issue_tasks=sum(row.verification_issue for row in rows),
        total_photos=sum(row.photo_count for row in rows),
        required_slots=required_slots,
        collected_required_slots=collected_slots,
    )


_COLORS = {
    "ink": "171717",
    "white": "FFFFFF",
    "subtle": "F5F5F5",
    "border": "D9D9D9",
    "muted": "6B6B6B",
    "green": "0F7B43",
    "green_soft": "EAF6EF",
    "yellow": "946200",
    "yellow_soft": "FFF4CC",
    "red": "B42318",
    "red_soft": "FDECEA",
}


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=_COLORS[color])


_THIN_BORDER = Border(
    left=Side(style="thin", color=_COLORS["border"]),
    right=Side(style="thin", color=_COLORS["border"]),
    top=Side(style="thin", color=_COLORS["border"]),
    bottom=Side(style="thin", color=_COLORS["border"]),
)


def _style_header_row(sheet, row: int, first_col: int, last_col: int) -> None:
    for col in range(first_col, last_col + 1):
        cell = sheet.cell(row=row, column=col)
        cell.fill = _fill("ink")
        cell.font = Font(color=_COLORS["white"], bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _THIN_BORDER
    sheet.row_dimensions[row].height = 28


def _set_widths(sheet, widths: Mapping[int, float]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width


def _photo_fill(row: ReportRow) -> PatternFill:
    return _fill("green_soft" if row.photo_complete else "red_soft")


def _parse_updated_at(value: str):
    try:
        return datetime.strptime(value, "%Y-%m-%d %H:%M:%S")
    except (TypeError, ValueError):
        return value


def _build_dashboard_sheet(workbook: Workbook, report: CollectionReport, work_root: str, generated_at: datetime) -> None:
    sheet = workbook.active
    sheet.title = "采集看板"
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:G2")
    title = sheet["A1"]
    title.value = "PhotoCollect 采集看板"
    title.fill = _fill("ink")
    title.font = Font(color=_COLORS["white"], bold=True, size=20)
    title.alignment = Alignment(horizontal="left", vertical="center")
    for row in sheet["A1:G2"]:
        for cell in row:
            cell.fill = _fill("ink")
    sheet.row_dimensions[1].height = 26
    sheet.row_dimensions[2].height = 26

    sheet.merge_cells("A3:G3")
    sheet["A3"] = f"工作目录：{work_root}    导出时间：{generated_at:%Y-%m-%d %H:%M}"
    sheet["A3"].font = Font(color=_COLORS["muted"], size=10)
    sheet["A3"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[3].height = 22

    kpis = (
        ("A5", "总任务", "B5", report.total_tasks, "subtle"),
        ("C5", "照片齐全", "D5", report.photo_complete_tasks, "green_soft"),
        ("E5", "缺照片", "F5", report.missing_photo_tasks, "red_soft"),
        ("A6", "待核验/异常", "B6", report.verification_issue_tasks, "yellow_soft"),
        ("C6", "照片总数", "D6", report.total_photos, "subtle"),
        ("E6", "必填照片项完成率", "F6", report.photo_completion_rate, "green_soft"),
    )
    for label_ref, label, value_ref, value, fill_name in kpis:
        label_cell = sheet[label_ref]
        value_cell = sheet[value_ref]
        label_cell.value = label
        value_cell.value = value
        label_cell.fill = _fill(fill_name)
        value_cell.fill = _fill(fill_name)
        label_cell.font = Font(color=_COLORS["muted"], bold=True, size=10)
        value_cell.font = Font(color=_COLORS["ink"], bold=True, size=14)
        label_cell.alignment = value_cell.alignment = Alignment(vertical="center")
        label_cell.border = value_cell.border = _THIN_BORDER
    sheet["F6"].number_format = "0.0%"
    sheet["F6"] = "=IF(SUM(C10:C17)=0,0,SUM(D10:D17)/SUM(C10:C17))"
    sheet["G5"] = "照片齐全率"
    sheet["G5"].font = Font(color=_COLORS["muted"], bold=True, size=10)
    sheet["G5"].fill = _fill("subtle")
    sheet["G5"].border = _THIN_BORDER
    sheet["G6"] = "=IF(B5=0,0,D5/B5)"
    sheet["G6"].font = Font(color=_COLORS["ink"], bold=True, size=14)
    sheet["G6"].fill = _fill("subtle")
    sheet["G6"].border = _THIN_BORDER
    sheet["G6"].number_format = "0.0%"
    sheet["G5"].alignment = sheet["G6"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[5].height = 28
    sheet.row_dimensions[6].height = 30

    sheet.merge_cells("A8:G8")
    sheet["A8"] = "八类照片收集情况"
    sheet["A8"].font = Font(color=_COLORS["ink"], bold=True, size=13)
    sheet["A8"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[8].height = 25
    headers = ("编号", "照片分类", "必填箱数", "已收集箱数", "缺失箱数", "完成率", "已上传照片")
    for col, value in enumerate(headers, 1):
        sheet.cell(row=9, column=col, value=value)
    _style_header_row(sheet, 9, 1, 7)

    for offset, category in enumerate(report.categories, 10):
        values = (
            category.index,
            category.label,
            category.required_tasks,
            category.collected_tasks,
            None,
            None,
            category.uploaded_photos,
        )
        for col, value in enumerate(values, 1):
            cell = sheet.cell(row=offset, column=col, value=value)
            cell.border = _THIN_BORDER
            cell.alignment = Alignment(
                horizontal="left" if col == 2 else "center",
                vertical="center",
                wrap_text=True,
            )
        sheet.cell(offset, 5, f"=C{offset}-D{offset}")
        sheet.cell(offset, 6, f'=IF(C{offset}=0,"-",D{offset}/C{offset})')
        for column in (5, 6):
            sheet.cell(offset, column).border = _THIN_BORDER
            sheet.cell(offset, column).alignment = Alignment(horizontal="center", vertical="center")
        sheet.cell(offset, 6).number_format = "0.0%"
        if not category.required_tasks:
            for column in (5, 6):
                sheet.cell(offset, column).fill = _fill("subtle")
                sheet.cell(offset, column).font = Font(color=_COLORS["muted"])
        elif category.missing_tasks:
            sheet.cell(offset, 5).fill = _fill("red_soft")
            sheet.cell(offset, 5).font = Font(color=_COLORS["red"], bold=True)
        else:
            sheet.cell(offset, 5).fill = _fill("green_soft")
            sheet.cell(offset, 5).font = Font(color=_COLORS["green"], bold=True)
        sheet.row_dimensions[offset].height = 24

    note_row = 19
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=7)
    sheet.cell(note_row, 1, "说明：缺失统计按当前全局及单箱验收标准计算；非必填照片不计为缺失。")
    sheet.cell(note_row, 1).font = Font(color=_COLORS["muted"], italic=True, size=9)
    sheet.cell(note_row, 1).alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[note_row].height = 28
    sheet.freeze_panes = "A10"
    sheet.auto_filter.ref = "A9:G17"
    _set_widths(sheet, {1: 9, 2: 26, 3: 13, 4: 14, 5: 13, 6: 12, 7: 15})
    sheet.print_area = "A1:G19"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _build_task_sheet(workbook: Workbook, report: CollectionReport) -> None:
    sheet = workbook.create_sheet("任务明细")
    sheet.sheet_view.showGridLines = False
    headers = ("箱号", "箱型", "照片状态", "OCR/箱号核验", "照片数", "缺失数量", "缺失照片", "最后更新时间")
    sheet.append(headers)
    _style_header_row(sheet, 1, 1, len(headers))

    for row_index, row in enumerate(report.rows, 2):
        missing_text = "、".join(row.missing_labels)
        if not missing_text:
            missing_text = "无"
        sheet.append((
            row.name,
            row.size_type,
            "照片齐全" if row.photo_complete else "缺照片",
            OCR_STATUS_LABELS.get(row.ocr_status, row.ocr_status or "待识别"),
            row.photo_count,
            len(row.missing_category_ids),
            missing_text,
            _parse_updated_at(row.last_updated),
        ))
        for col in range(1, len(headers) + 1):
            cell = sheet.cell(row_index, col)
            cell.border = _THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=col in {4, 7})
        sheet.cell(row_index, 3).fill = _photo_fill(row)
        sheet.cell(row_index, 3).font = Font(bold=True)
        if row.verification_issue:
            sheet.cell(row_index, 4).fill = _fill("yellow_soft")
            sheet.cell(row_index, 4).font = Font(color=_COLORS["yellow"], bold=True)
        if row.missing_category_ids:
            sheet.cell(row_index, 6).fill = _fill("red_soft")
            sheet.cell(row_index, 7).fill = _fill("red_soft")
            sheet.cell(row_index, 7).font = Font(color=_COLORS["red"], bold=True)
        if isinstance(sheet.cell(row_index, 8).value, datetime):
            sheet.cell(row_index, 8).number_format = "yyyy-mm-dd hh:mm"
        sheet.row_dimensions[row_index].height = 30

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:H{max(1, len(report.rows) + 1)}"
    _set_widths(sheet, {1: 18, 2: 11, 3: 13, 4: 20, 5: 10, 6: 11, 7: 48, 8: 20})
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _build_matrix_sheet(workbook: Workbook, report: CollectionReport) -> None:
    sheet = workbook.create_sheet("照片矩阵")
    sheet.sheet_view.showGridLines = False
    headers = ["箱号", "箱型", "照片状态"]
    headers.extend(f"{category.index} {category.label}" for category in report.categories)
    headers.append("缺失汇总")
    sheet.append(headers)
    _style_header_row(sheet, 1, 1, len(headers))

    for row_index, row in enumerate(report.rows, 2):
        values = [
            row.name,
            row.size_type,
            "照片齐全" if row.photo_complete else "缺照片",
        ]
        for category in report.categories:
            count = row.category_count(category.index)
            if category.index in row.required_categories:
                values.append(f"已收 {count}张" if count else "缺失")
            else:
                values.append(f"选填 {count}张" if count else "-")
        values.append("、".join(row.missing_labels) or "无")
        sheet.append(values)

        for col in range(1, len(headers) + 1):
            cell = sheet.cell(row_index, col)
            cell.border = _THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.cell(row_index, 3).fill = _photo_fill(row)
        for offset, category in enumerate(report.categories, 4):
            cell = sheet.cell(row_index, offset)
            count = row.category_count(category.index)
            if category.index in row.required_categories and not count:
                cell.fill = _fill("red_soft")
                cell.font = Font(color=_COLORS["red"], bold=True)
            elif count:
                cell.fill = _fill("green_soft")
                cell.font = Font(color=_COLORS["green"])
            else:
                cell.fill = _fill("subtle")
                cell.font = Font(color=_COLORS["muted"])
        sheet.row_dimensions[row_index].height = 30

    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = f"A1:L{max(1, len(report.rows) + 1)}"
    widths = {1: 18, 2: 11, 3: 13, 12: 42}
    widths.update({column: 18 for column in range(4, 12)})
    _set_widths(sheet, widths)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def export_collection_report_xlsx(
    report: CollectionReport,
    destination: Path | str,
    work_root: Path | str,
    generated_at: datetime | None = None,
    cancel_event=None,
    progress_callback=None,
) -> Path:
    """Write a three-sheet Excel report and publish it atomically."""
    output_path = Path(destination)
    if output_path.suffix.lower() != ".xlsx":
        output_path = output_path.with_suffix(".xlsx")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    timestamp = generated_at or datetime.now()
    partial_path = output_path.with_name(f".{output_path.stem}.{uuid.uuid4().hex}.partial.xlsx")

    def check_cancelled() -> None:
        if cancel_event is not None and cancel_event.is_set():
            raise ReportExportCancelled()

    def report_progress(current: int, label: str) -> None:
        if progress_callback is not None:
            progress_callback(current, 5, label)

    workbook = Workbook()
    try:
        check_cancelled()
        workbook.properties.creator = "RenYiZheng"
        workbook.properties.title = "PhotoCollect 采集看板"
        workbook.properties.subject = "集装箱照片收集情况与缺项明细"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"

        report_progress(1, "生成采集概览")
        _build_dashboard_sheet(workbook, report, str(work_root), timestamp)
        check_cancelled()
        report_progress(2, "生成任务明细")
        _build_task_sheet(workbook, report)
        check_cancelled()
        report_progress(3, "生成照片矩阵")
        _build_matrix_sheet(workbook, report)
        check_cancelled()
        report_progress(4, "写入 Excel 文件")
        workbook.save(partial_path)
        check_cancelled()
        report_progress(5, "校验并发布")
        with zipfile.ZipFile(partial_path, "r") as archive:
            if archive.testzip() is not None or "xl/workbook.xml" not in archive.namelist():
                raise ValueError("生成的Excel文件结构无效")
        check_cancelled()
        os.replace(partial_path, output_path)
    finally:
        workbook.close()
        if partial_path.exists():
            partial_path.unlink()
    return output_path
