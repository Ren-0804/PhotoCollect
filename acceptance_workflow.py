"""Repeatable isolated acceptance workflow for source and packaged releases."""

from __future__ import annotations

import csv
import io
import json
import os
import shutil
import sys
import tempfile
import time
import traceback
import zipfile
from datetime import datetime
from pathlib import Path

from health_services import scan_data_health
from report_services import export_collection_report_xlsx
from storage_services import (
    TaskExportRecord,
    TrashStore,
    export_handoff_package,
    file_sha256,
)


def _atomic_write_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    partial = path.with_name(f"{path.name}.partial")
    try:
        partial.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        os.replace(partial, path)
    finally:
        if partial.exists():
            partial.unlink()


def _capture_clipboard(app_module):
    clipboard = app_module.QApplication.clipboard()
    mime = clipboard.mimeData()
    if mime is None:
        return {}
    return {
        str(fmt): bytes(mime.data(fmt))
        for fmt in mime.formats()
    }


def _restore_clipboard(app_module, formats: dict[str, bytes]) -> None:
    if not formats:
        app_module.QApplication.clipboard().clear()
        return
    mime = app_module.QMimeData()
    for fmt, payload in formats.items():
        mime.setData(fmt, app_module.QByteArray(payload))
    app_module.QApplication.clipboard().setMimeData(mime)


def run_acceptance(
    app_module,
    report_path: Path | str,
    *,
    workspace: Path | str | None = None,
    ocr_fixture: Path | str | None = None,
) -> int:
    """Run the real local workflow in an isolated directory and write JSON."""
    report_path = Path(report_path)
    cleanup_workspace = workspace is None
    if cleanup_workspace:
        report_path.parent.mkdir(parents=True, exist_ok=True)
        workspace = Path(tempfile.mkdtemp(
            prefix=f".{report_path.stem}_workspace_",
            dir=report_path.parent,
        ))
    workspace = Path(workspace)
    if workspace.exists() and any(workspace.iterdir()):
        payload = {
            "ok": False,
            "version": str(getattr(app_module, "APP_VERSION", "")),
            "error": f"验收工作目录不是空目录：{workspace}",
            "steps": [],
        }
        _atomic_write_json(report_path, payload)
        return 1
    workspace.mkdir(parents=True, exist_ok=True)

    root = workspace / "collection"
    inputs = workspace / "inputs"
    artifacts = workspace / "artifacts"
    root.mkdir()
    inputs.mkdir()
    artifacts.mkdir()

    application = app_module.QApplication.instance() or app_module.QApplication([])
    original_config_path = app_module.CONFIG_PATH
    clipboard_formats = _capture_clipboard(app_module)
    steps: list[dict] = []
    result: dict = {
        "ok": False,
        "version": str(getattr(app_module, "APP_VERSION", "")),
        "started_at": datetime.now().isoformat(timespec="seconds"),
        "workspace": str(workspace),
        "steps": steps,
    }
    runtime_path = Path(sys.executable)
    result["runtime"] = {
        "frozen": bool(getattr(sys, "frozen", False)),
        "executable": str(runtime_path),
        "size_bytes": runtime_path.stat().st_size if runtime_path.is_file() else 0,
        "sha256": file_sha256(runtime_path) if runtime_path.is_file() else "",
    }
    state: dict = {}
    window = None
    host = None

    def run_step(name: str, operation):
        started = time.perf_counter()
        try:
            detail = operation()
            steps.append({
                "name": name,
                "ok": True,
                "duration_ms": round((time.perf_counter() - started) * 1000),
                "detail": detail or "",
            })
            return detail
        except Exception as exc:
            steps.append({
                "name": name,
                "ok": False,
                "duration_ms": round((time.perf_counter() - started) * 1000),
                "error": str(exc),
            })
            raise

    def require(condition: bool, message: str) -> None:
        if not condition:
            raise AssertionError(message)

    try:
        app_module.CONFIG_PATH = workspace / "app-config.json"
        complete = "MSCU6639870"
        pending = "LYGU3043781"
        anomaly = "BAD-TASK"
        empty_prefix = "TGHU123456"
        empty = empty_prefix + str(app_module.iso6346_check_digit(empty_prefix))
        task_names = (complete, pending, anomaly, empty)
        state.update({
            "complete": complete,
            "pending": pending,
            "anomaly": anomaly,
            "empty": empty,
        })

        def batch_create():
            analysis = app_module.analyze_batch_input("\n".join(task_names), set())
            require(len(analysis.creatable) == 4, "批量创建预览数量不正确")
            require(analysis.valid_count == 3, "有效箱号数量不正确")
            require(analysis.warning_count == 1, "异常箱号数量不正确")
            for candidate in analysis.creatable:
                (root / candidate.name).mkdir()
            app_module.write_meta(
                root / app_module.META_NAME,
                {"__requirements__": list(range(1, 9))},
            )
            return {"created": len(analysis.creatable), "warning": analysis.warning_count}

        run_step("batch_create", batch_create)

        host = app_module.QWidget()
        host.conf = {"auto_ocr_enabled": False}
        notices = []
        host.show_toast = lambda *args: notices.append(args)
        state["notices"] = notices
        upload_state = {}

        def local_upload():
            detail = app_module.BoxDetailWindow(root, complete, host)
            image = app_module.QImage(96, 72, app_module.QImage.Format_RGB32)
            image.fill(app_module.QColor(210, 40, 40))
            source = inputs / "local.png"
            require(image.save(str(source), "PNG"), "无法生成本地上传测试图")
            detail.add_files(1, [str(source)])
            require(len(detail.meta["1"]) == 1, "本地文件未保存到分类 1")
            upload_state["detail"] = detail
            upload_state["local_source"] = source
            return {"category": 1, "files": len(detail.meta["1"])}

        run_step("local_upload", local_upload)

        def telegram_virtual_upload():
            detail = upload_state["detail"]
            image = app_module.QImage(96, 72, app_module.QImage.Format_RGB32)
            image.fill(app_module.QColor(40, 160, 80))
            payload = detail._qimage_png_bytes(image)
            require(bool(payload), "无法生成 Telegram 虚拟文件测试数据")
            detail.add_virtual_files(2, [("telegram.png", payload)])
            require(len(detail.meta["2"]) == 1, "Telegram 虚拟文件未保存到分类 2")
            return {"category": 2, "files": len(detail.meta["2"])}

        run_step("telegram_virtual_upload", telegram_virtual_upload)

        def clipboard_upload():
            detail = upload_state["detail"]
            image = app_module.QImage(96, 72, app_module.QImage.Format_RGB32)
            image.fill(app_module.QColor(50, 90, 210))
            app_module.QApplication.clipboard().setImage(image)
            detail.select_category(3)
            detail.paste_to_selected_category()
            require(len(detail.meta["3"]) == 1, "剪贴板图片未保存到分类 3")
            return {"category": 3, "files": len(detail.meta["3"])}

        run_step("clipboard_upload", clipboard_upload)

        def complete_remaining_categories():
            detail = upload_state["detail"]
            for category in range(4, 9):
                image = app_module.QImage(96, 72, app_module.QImage.Format_RGB32)
                image.fill(app_module.QColor(
                    20 + category * 20,
                    30 + category * 13,
                    40 + category * 11,
                ))
                detail.add_image_data(category, image)
                require(len(detail.meta[str(category)]) == 1, f"分类 {category} 未保存")
            detail.prepare_for_disposal()
            detail.close()

            pending_detail = app_module.BoxDetailWindow(root, pending, host)
            image = app_module.QImage(96, 72, app_module.QImage.Format_RGB32)
            image.fill(app_module.QColor(130, 60, 180))
            pending_detail.add_image_data(1, image)
            require(len(pending_detail.meta["1"]) == 1, "待补任务图片未保存")
            pending_detail.prepare_for_disposal()
            pending_detail.close()
            return {"complete_photos": 8, "pending_photos": 1}

        run_step("complete_and_pending_states", complete_remaining_categories)

        def duplicate_detection():
            detail = app_module.BoxDetailWindow(root, complete, host)
            before = sum(len(detail.meta[str(index)]) for index in range(1, 9))
            detail.add_files(8, [str(upload_state["local_source"])])
            after = sum(len(detail.meta[str(index)]) for index in range(1, 9))
            detail.prepare_for_disposal()
            detail.close()
            require(before == after == 8, "重复图片没有被跳过")
            return {"before": before, "after": after}

        run_step("duplicate_detection", duplicate_detection)

        def delete_and_restore():
            meta = app_module.read_meta(root / app_module.META_NAME)
            filename = meta[complete]["8"][0]
            store = TrashStore(root)
            photo_entry = store.trash_photo(complete, 8, filename)
            require(not (root / complete / filename).exists(), "照片没有移入最近删除")
            store.restore(photo_entry["id"])
            require((root / complete / filename).is_file(), "照片没有恢复")

            task_entry = store.trash_task(empty)
            require(not (root / empty).exists(), "任务没有移入最近删除")
            TrashStore(root).restore(task_entry["id"])
            require((root / empty).is_dir(), "任务没有从新的回收区实例恢复")
            return {"photo_restore": True, "task_restore": True}

        run_step("delete_and_restore", delete_and_restore)

        def dashboard_roundtrip():
            nonlocal window
            app_module.CONFIG_PATH.write_text(json.dumps({
                "root": str(root),
                "hist": [str(root)],
                "dashboard_filter": "all",
                "sort_order": "status",
                "sidebar_collapsed": False,
                "auto_ocr_enabled": False,
            }, ensure_ascii=False), encoding="utf-8")
            window = app_module.ContainerApp()
            require(len(window.task_data) == 4, "首页任务数量不正确")
            started = time.perf_counter()
            window.open_box_name(pending)
            require(window._detail_page is not None, "无法打开任务详情")
            window.on_back_to_dashboard()
            application.processEvents()
            elapsed_ms = round((time.perf_counter() - started) * 1000)
            require(window.stack.currentWidget() is window.dashboard_page, "未返回任务工作台")
            require(elapsed_ms < 1000, f"详情往返耗时过长：{elapsed_ms} ms")
            return {"tasks": len(window.task_data), "roundtrip_ms": elapsed_ms}

        run_step("dashboard_roundtrip", dashboard_roundtrip)

        def ocr_switch():
            window.set_auto_ocr_enabled(True)
            require(window.conf["auto_ocr_enabled"] is True, "自动 OCR 未开启")
            window.set_auto_ocr_enabled(False)
            require(window.conf["auto_ocr_enabled"] is False, "自动 OCR 未关闭")
            complete_summary = next(
                item for item in window.task_data
                if item["name"] == complete
            )
            require(complete_summary["task_state"] == "complete", "关闭自动 OCR 后完整任务状态错误")
            return {"enabled_then_disabled": True}

        run_step("ocr_switch", ocr_switch)

        def responsive_layout():
            checked = []
            for width, height in ((1024, 700), (1366, 768), (1920, 1080)):
                window.resize(width, height)
                window._apply_command_bar_layout()
                window.centralWidget().layout().activate()
                application.processEvents()
                require(window.btn_close.width() > 0, "关闭按钮不可用")
                require(window.btn_maximize.width() > 0, "最大化按钮不可用")
                require(window.btn_minimize.width() > 0, "最小化按钮不可用")
                require(window.command_bar.sizeHint().height() > 0, "命令栏布局无效")
                checked.append(f"{width}x{height}")
            require(window.minimumWidth() <= 400, "窗口最小宽度不支持窄分屏")
            return {"viewports": checked, "minimum_width": window.minimumWidth()}

        run_step("responsive_layout", responsive_layout)

        def zip_export():
            selected = list(window.task_data)
            state_labels = {"complete": "已完成", "pending": "待补", "anomaly": "异常"}
            records = [
                TaskExportRecord(
                    name=item["name"],
                    size_type=item.get("ocr_size_type", ""),
                    ocr_status=item.get("ocr_status", ""),
                    task_status=state_labels.get(item.get("task_state"), "待补"),
                    photo_count=item.get("pic_cnt", 0),
                    missing=tuple(item.get("missing", [])),
                    last_updated=item.get("last_updated", ""),
                    required_categories=tuple(item.get("required_categories", [])),
                    missing_category_ids=tuple(item.get("missing_category_ids", [])),
                )
                for item in selected
            ]
            destination = artifacts / "handoff.zip"
            summary = export_handoff_package(
                root,
                destination,
                records,
                dict(app_module.ALL_CATEGORIES),
                image_transformer=app_module.optimize_export_image,
            )
            with zipfile.ZipFile(destination) as archive:
                names = set(archive.namelist())
                require("summary.csv" in names and "data.json" in names, "ZIP 缺少清单文件")
                require(not any(name.startswith("containers/") for name in names), "ZIP 出现 containers 包装目录")
                photo_names = {
                    name for name in names
                    if not name.endswith("/") and name not in {"summary.csv", "data.json"}
                }
                flat_layout = all(name.count("/") == 1 for name in photo_names)
                require(flat_layout, "箱号目录下出现额外子目录")
                exported = json.loads(archive.read("data.json").decode("utf-8"))
                manifest_names = {
                    f"{record.name}/{filename}"
                    for record in records
                    if record.name in exported
                    for category in map(str, range(1, 9))
                    for filename in exported[record.name][category]
                }
                manifest_matches = manifest_names == photo_names
                require(manifest_matches, "data.json 照片路径与 ZIP 不一致")
                csv_rows = list(csv.DictReader(io.StringIO(
                    archive.read("summary.csv").decode("utf-8-sig")
                )))
                require(len(csv_rows) == summary.task_count, "CSV 任务数量不正确")
                require(f"{empty}/" in names, "空任务箱号目录没有保留")
            result["zip_validation"] = {
                "flat_container_layout": flat_layout,
                "manifest_matches_files": manifest_matches,
                "task_count": summary.task_count,
                "photo_count": summary.photo_count,
            }
            return result["zip_validation"]

        run_step("zip_export", zip_export)

        def excel_export():
            report = app_module.build_collection_report(
                window.task_data,
                app_module.ALL_CATEGORIES,
            )
            destination = artifacts / "collection.xlsx"
            export_collection_report_xlsx(report, destination, root)
            from openpyxl import load_workbook

            workbook = load_workbook(destination, read_only=True, data_only=False)
            try:
                require(
                    workbook.sheetnames == ["采集看板", "任务明细", "照片矩阵"],
                    "Excel 工作表结构不正确",
                )
            finally:
                workbook.close()
            require(report.total_tasks == 4, "Excel 报表任务数量不正确")
            return {"tasks": report.total_tasks, "sheets": 3}

        run_step("excel_export", excel_export)

        def health_check():
            report = scan_data_health(
                root,
                task_name_validator=lambda name: bool(
                    app_module.extract_declared_container_no(name)
                    and app_module.validate_container_no(
                        app_module.extract_declared_container_no(name)
                    )
                ),
            )
            codes = {issue.code for issue in report.issues}
            require("invalid_task_name" in codes, "数据检查未识别异常任务名")
            require(report.error_count == 0, "隔离工作目录出现非预期数据错误")
            return {"warnings": report.warning_count, "errors": report.error_count}

        run_step("data_health", health_check)

        def ocr_regression():
            require(ocr_fixture is not None, "OCR固定样图未提供")
            fixture = Path(ocr_fixture)
            require(fixture.is_file(), f"OCR 样图不存在：{fixture}")
            ocr_result = app_module.run_ocr_image(str(fixture))
            expected_container = getattr(
                app_module,
                "OCR_REGRESSION_CONTAINER_NO",
                "LYGU3043781",
            )
            expected_size_type = getattr(
                app_module,
                "OCR_REGRESSION_SIZE_TYPE",
                "22G1",
            )
            require(
                ocr_result.get("container_no") == expected_container,
                "OCR 箱号回归失败",
            )
            require(
                ocr_result.get("size_type") == expected_size_type,
                "OCR 箱型回归失败",
            )
            require(
                bool(ocr_result.get("check_digit_valid")),
                "OCR 箱号校验位回归失败",
            )
            return {
                "container_no": ocr_result.get("container_no"),
                "size_type": ocr_result.get("size_type"),
                "check_digit_valid": bool(ocr_result.get("check_digit_valid")),
                "elapsed_ms": ocr_result.get("elapsed_ms", 0),
            }

        run_step("ocr_fixture", ocr_regression)

        def photo_classifier_regression():
            # The model is git-ignored, so a source checkout may lack it; the
            # packaged EXE must always carry a working one.
            classifier = app_module.get_photo_classifier()
            if classifier is None and not getattr(sys, "frozen", False):
                return {"skipped": "源码目录没有 models/photo_classifier.onnx"}
            require(classifier is not None, "自动分类模型没有打包进程序")
            prediction = classifier.predict(str(ocr_fixture))
            require(prediction.category in (4, 5), f"自动分类回归失败：关门样图被判为 {prediction.category}")
            return {"category": prediction.category, "confidence": round(prediction.confidence, 4)}

        run_step("photo_classifier", photo_classifier_regression)
        result["ok"] = True
    except Exception as exc:
        result["error"] = str(exc)
        result["traceback"] = traceback.format_exc()
    finally:
        if window is not None:
            window.close()
        if host is not None:
            host.close()
        app_module.QThreadPool.globalInstance().waitForDone(5000)
        _restore_clipboard(app_module, clipboard_formats)
        app_module.CONFIG_PATH = original_config_path
        result["completed_at"] = datetime.now().isoformat(timespec="seconds")
        result["duration_ms"] = sum(
            int(step.get("duration_ms", 0))
            for step in steps
        )
        if cleanup_workspace:
            try:
                shutil.rmtree(workspace)
            except OSError as exc:
                result["ok"] = False
                result["error"] = f"验收工作目录清理失败：{exc}"
                result["cleanup_error"] = str(exc)
        _atomic_write_json(report_path, result)
    return 0 if result["ok"] else 1
